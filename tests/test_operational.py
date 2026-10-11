import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from ats_score import calculate_ats_score
from match_score import assess_eligibility, calculate_match_score, identify_skill_requirements
from profile_support import configured_path, load_profile_config, safe_output_path
from search_linkedin import filter_by_keywords
from validate_tailoring import validate_entry


SDET = {
    "skills": ["java", "selenium", "junit"],
    "domains": ["test automation"],
    "preferred_roles": ["sdet", "test automation"],
    "locations": {"remote_us": True, "local": ["nashville", "tennessee"], "home_state": "tennessee", "no_relocation": True},
    "employment_history": [{"employer": "Quality Labs"}],
}
SWE = {
    "skills": ["java", "spring boot", "kafka"],
    "domains": ["backend"],
    "preferred_roles": ["java software engineer", "backend engineer"],
    "locations": {"remote_us": True, "local": ["chicago"]},
    "employment_history": [{"employer": "Software Works"}],
}


class FilteringTests(unittest.TestCase):
    def test_sdet_and_swe_titles_are_profile_driven(self):
        jobs = [
            {"title": "Senior SDET", "url": "https://example/1"},
            {"title": "Senior Java Software Engineer", "url": "https://example/2"},
            {"title": "Engineering Manager", "url": "https://example/3"},
            {"title": "", "url": "https://example/4"},
        ]
        sdet = {"search": {"title_include": ["sdet", "test automation"], "title_exclude": ["manager"]}}
        swe = {"search": {"title_include": ["java", "backend"], "title_exclude": ["manager"]}}
        self.assertEqual([jobs[0]], filter_by_keywords(jobs, "", sdet))
        self.assertEqual([jobs[1]], filter_by_keywords(jobs, "", swe))


class ProfileAndScoringTests(unittest.TestCase):
    def test_profile_paths_are_isolated(self):
        with tempfile.TemporaryDirectory() as first, tempfile.TemporaryDirectory() as second:
            for directory, name in ((first, "one"), (second, "two")):
                Path(directory, "config.json").write_text(json.dumps({
                    "candidate": {"name": name},
                    "paths": {"journal": "data/journal.xlsx"}
                }), encoding="utf-8")
            dir_one, config_one = load_profile_config(first)
            dir_two, config_two = load_profile_config(second)
            self.assertNotEqual(
                configured_path(dir_one, config_one, "journal", "journal.xlsx"),
                configured_path(dir_two, config_two, "journal", "journal.xlsx"),
            )

    def test_candidate_specific_scoring(self):
        description = (
            "We need Java Selenium JUnit test automation experience for a remote "
            "Senior SDET building reliable automated test suites and quality systems."
        )
        job = {"title": "Senior SDET", "location": "Remote USA"}
        config = {"scoring": {"skill_vocabulary": ["java", "selenium", "junit", "spring boot", "kafka"]}}
        sdet_score, sdet_details = calculate_match_score(job, description, "", SDET, config)
        swe_score, swe_details = calculate_match_score(job, description, "", SWE, config)
        self.assertGreater(sdet_score, swe_score)
        self.assertIn("selenium", sdet_details["verified_matches"])
        self.assertIn("selenium", swe_details["missing_qualifications"])

    def test_missing_description_is_unknown_and_retryable(self):
        score, details = calculate_match_score({"title": "SDET"}, "", "", SDET, {})
        ats, ats_details = calculate_ats_score("", "", SDET)
        self.assertIsNone(score)
        self.assertEqual("Unknown", details["fit_indicator"])
        self.assertIsNone(ats)
        self.assertIn("unavailable", ats_details["reason"].lower())


class EligibilityRegressionTests(unittest.TestCase):
    def classify(self, location, description):
        return assess_eligibility({"location": location}, description, SDET)

    def test_el01_cognizant_hybrid_raleigh_is_ineligible_and_not_actionable(self):
        description = "This role requires Selenium and C#. Hybrid attendance is required three days per week in Raleigh."
        score, details = calculate_match_score(
            {"title": "QA Automation Engineer", "location": "Raleigh, NC"},
            description, "", SDET, {},
        )
        self.assertEqual("Ineligible", details["eligibility"])
        self.assertIn("c#", details["missing_qualifications"])
        from run_scoring import determine_priority
        self.assertEqual("Ineligible", determine_priority(score, details["eligibility"]))

    def test_el02_confirmed_remote_us_is_eligible(self):
        state, _ = self.classify("Remote USA", "This is a fully remote role available throughout the United States, including Tennessee.")
        self.assertEqual("Eligible", state)

    def test_el03_nashville_onsite_is_eligible(self):
        state, _ = self.classify("Nashville, TN", "Employees work on-site in our Nashville office.")
        self.assertEqual("Eligible", state)

    def test_el04_distant_unspecified_is_unverified(self):
        state, _ = self.classify("Denver, CO", "Join our quality engineering organization and improve test systems.")
        self.assertEqual("Unverified", state)

    def test_el05_united_states_without_remote_evidence_is_unverified(self):
        state, _ = self.classify("United States", "The team builds enterprise quality systems.")
        self.assertEqual("Unverified", state)

    def test_el06_required_relocation_is_ineligible(self):
        state, _ = self.classify("Seattle, WA", "The selected candidate must relocate to Seattle before starting.")
        self.assertEqual("Ineligible", state)

    def test_el07_remote_excluding_home_state_is_ineligible(self):
        state, _ = self.classify("Remote USA", "This fully remote role is available nationwide except in Tennessee.")
        self.assertEqual("Ineligible", state)

    def test_el08_conflicting_remote_and_hybrid_is_unverified(self):
        state, _ = self.classify("Remote USA", "The position is fully remote with a flexible hybrid option.")
        self.assertEqual("Unverified", state)

    def test_existing_profile_without_new_location_fields_remains_compatible(self):
        legacy_candidate = {"locations": {"remote_us": True, "local": ["nashville"]}}
        state, _ = assess_eligibility(
            {"location": "Remote USA"}, "This is a fully remote role within the United States.", legacy_candidate
        )
        self.assertEqual("Eligible", state)


class RequirementRegressionTests(unittest.TestCase):
    def test_sc01_required_csharp_is_missing_for_java_candidate(self):
        description = "Candidates must have Selenium and C# experience for this automation position."
        score, details = calculate_match_score(
            {"title": "SDET", "location": "Remote USA"}, description, "", SDET, {}
        )
        self.assertIn("c#", details["identified_required_skills"])
        self.assertIn("c#", details["missing_qualifications"])
        self.assertIn("selenium", details["verified_skill_matches"])

    def test_sc02_preferred_skill_is_not_mandatory(self):
        requirements = identify_skill_requirements("Java is required. C# is preferred.", {}, SDET)
        self.assertIn("java", requirements["required"])
        self.assertIn("c#", requirements["preferred"])
        self.assertNotIn("c#", requirements["required"])

    def test_sc03_embedded_skill_name_is_not_false_positive(self):
        requirements = identify_skill_requirements("JavaScript is required for browser development.", {}, SDET)
        self.assertNotIn("java", requirements["required"])
        self.assertIn("javascript", requirements["required"])

    def test_sc04_structured_requirements_are_respected(self):
        job = {"title": "SDET", "location": "Remote USA", "required_skills": ["C#", "Selenium"]}
        _, details = calculate_match_score(job, "A detailed remote quality engineering position with many responsibilities.", "", SDET, {})
        self.assertEqual(["c#", "selenium"], details["identified_required_skills"])
        self.assertIn("c#", details["missing_qualifications"])

    def test_sc05_unspecified_requirements_are_unknown(self):
        description = "Our engineering organization uses Java and Selenium across several teams and products."
        _, details = calculate_match_score({"title": "SDET", "location": "Remote USA"}, description, "", SDET, {})
        self.assertEqual("unknown", details["component_evidence"]["skills"]["status"])
        self.assertIn("ambiguous_skill_requirements", details["unknown"])
        self.assertEqual(0, details["components"]["skills"])

    def test_sc06_strong_required_skill_matches_are_positive_evidence(self):
        description = "This remote role requires Java, Selenium, and JUnit experience building test automation suites."
        score, details = calculate_match_score({"title": "Senior SDET", "location": "Remote USA"}, description, "", SDET, {})
        self.assertGreaterEqual(score, 75)
        self.assertEqual(["java", "junit", "selenium"], details["verified_skill_matches"])


class ValidationAndSafetyTests(unittest.TestCase):
    def _entry(self):
        return {
            "url": "https://example/job",
            "company": "Example",
            "title": "Senior SDET",
            "tailored_summary": "Automation engineer. Builds reliable tests.",
            "tailored_strengths": ["Automation"],
            "tailored_highlights": [{
                "header": "Quality Labs | Senior SDET | 2021–Present",
                "bullets": ["Built Java Selenium suites that improved regression feedback through CI."],
            }],
            "ats_keywords_injected": ["java", "selenium"],
            "claimed_qualifications": ["java"],
            "cover_letter": "Dear Hiring Manager,\n\nI build verified automation systems.\n\nSincerely,\nCandidate",
        }

    def test_resume_validation_rejects_fabricated_qualification(self):
        entry = self._entry()
        entry["claimed_qualifications"].append("cissp")
        errors = validate_entry(entry, candidate=SDET, settings={"require_supported_keywords": True})
        self.assertTrue(any("unsupported claimed qualifications" in error for error in errors))

    def test_output_overwrite_protection(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "resume.pdf")
            path.write_text("existing", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                safe_output_path(path)


class JournalTests(unittest.TestCase):
    def test_journal_integrity_dedup_and_init_protection(self):
        try:
            from journal import add_jobs, init_journal
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "journal.xlsx")
            init_journal(path)
            with self.assertRaises(FileExistsError):
                init_journal(path)
            jobs = [{"title": "Senior SDET", "company": "Acme", "location": "Remote", "url": "https://example/jobs/123"}]
            self.assertEqual((1, 0), add_jobs(jobs, path))
            self.assertEqual((0, 1), add_jobs(jobs, path))
            workbook = load_workbook(path)
            self.assertEqual([], __import__("journal").validate_journal_rows(workbook["Jobs"]))


class WorkflowTests(unittest.TestCase):
    def test_cognizant_case_is_written_as_ineligible_with_technical_score(self):
        try:
            from journal import add_jobs, init_journal
            from openpyxl import load_workbook
            import run_scoring
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            profile = Path(directory)
            config = {
                "candidate": SDET,
                "scoring": {
                    "priority_thresholds": {"high": 75, "medium": 50},
                    "skill_vocabulary": ["java", "selenium", "c#"],
                },
                "paths": {"journal": "journal.xlsx", "master_profile": "MASTER_PROFILE.md", "output": "output"},
            }
            (profile / "config.json").write_text(json.dumps(config), encoding="utf-8")
            (profile / "MASTER_PROFILE.md").write_text("Java and Selenium automation", encoding="utf-8")
            output = profile / "output"
            output.mkdir()
            journal = profile / "journal.xlsx"
            init_journal(journal)
            url = "https://example/jobs/cognizant-123"
            add_jobs([{"title": "QA Automation Engineer", "company": "Cognizant", "location": "Raleigh, NC", "url": url}], journal)
            description = "Applicants must have Selenium and C#. This hybrid role requires attendance three days per week in Raleigh."
            (output / "descriptions_cache.json").write_text(json.dumps({url: description}), encoding="utf-8")
            with patch.object(sys, "argv", ["run_scoring.py", "--profile-dir", str(profile)]):
                run_scoring.main()
            workbook = load_workbook(journal)
            row = list(workbook["Jobs"].iter_rows(min_row=2, values_only=True))[0]
            self.assertIsInstance(row[6], int)
            self.assertIn("c#", row[8])
            self.assertIn("Eligibility: Ineligible", row[11])
            self.assertEqual("Ineligible", row[12])

    def test_synthetic_profile_journal_to_scoring_workflow(self):
        try:
            from journal import add_jobs, init_journal
            from openpyxl import load_workbook
            import run_scoring
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            profile = Path(directory)
            config = {
                "candidate": SDET,
                "scoring": {
                    "priority_thresholds": {"high": 75, "medium": 50},
                    "skill_vocabulary": ["java", "selenium", "junit"],
                },
                "paths": {
                    "journal": "journal.xlsx",
                    "master_profile": "MASTER_PROFILE.md",
                    "output": "output",
                },
            }
            (profile / "config.json").write_text(json.dumps(config), encoding="utf-8")
            (profile / "MASTER_PROFILE.md").write_text("Java Selenium JUnit automation", encoding="utf-8")
            output = profile / "output"
            output.mkdir()
            journal = profile / "journal.xlsx"
            init_journal(journal)
            url = "https://example/jobs/987"
            add_jobs([{"title": "Senior SDET", "company": "Acme", "location": "Remote USA", "url": url}], journal)
            workbook = load_workbook(journal)
            workbook["Applications"].append(["2026-01-01", "Existing Co", "Existing Role", "2026-01-02", "v1", "No", "Applied"])
            workbook.save(journal)
            description = "Senior SDET role requiring Java Selenium and JUnit for reliable automated regression suites."
            (output / "descriptions_cache.json").write_text(json.dumps({url: description}), encoding="utf-8")
            with patch.object(sys, "argv", ["run_scoring.py", "--profile-dir", str(profile)]):
                run_scoring.main()
            with patch.object(sys, "argv", ["run_scoring.py", "--profile-dir", str(profile), "--rescore"]):
                run_scoring.main()
            workbook = load_workbook(journal)
            row = list(workbook["Jobs"].iter_rows(min_row=2, values_only=True))[0]
            self.assertIsInstance(row[6], int)
            self.assertIn(row[10], {"Strong", "Moderate", "Weak"})
            self.assertEqual(2, workbook["Jobs"].max_row)
            self.assertEqual("Existing Co", workbook["Applications"].cell(2, 2).value)
            self.assertEqual("Applied", workbook["Applications"].cell(2, 7).value)


if __name__ == "__main__":
    unittest.main()

